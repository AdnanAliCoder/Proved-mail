import csv
import io
import os
import re
from datetime import date, datetime
from pathlib import Path

EMAIL_FIELD_CANDIDATES = {"email", "emailaddress", "mail", "workemail", "businessemail", "e-mail"}

def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()

def detect_email_field(fieldnames):
    for field in fieldnames:
        key = re.sub(r"[-_\s]", "", str(field).strip().lower())
        if key in {"email", "emailaddress", "mail", "workemail", "businessemail"}:
            return field
    return None

def read_lead_file(file_path, original_filename):
    """Return (headers, rows, email_field); supports CSV and XLSX without pandas."""
    suffix = Path(original_filename).suffix.lower()
    if suffix == ".csv":
        with open(file_path, "rb") as stream:
            raw = stream.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError("The CSV file needs a header row.")
        headers = [str(h).strip() if h is not None else "" for h in reader.fieldnames]
        rows = []
        for row in reader:
            if not row or not any((v or "").strip() for v in row.values() if isinstance(v, str)):
                continue
            rows.append({headers[i]: _cell_text(row.get(reader.fieldnames[i])) for i in range(len(headers))})
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX support is unavailable; install openpyxl.") from exc
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            first = next(values, None)
            if not first:
                raise ValueError("The Excel sheet is empty.")
            headers = [(_cell_text(v) or f"Column {i + 1}") for i, v in enumerate(first)]
            rows = []
            for values_row in values:
                if not values_row or not any(v is not None and str(v).strip() for v in values_row):
                    continue
                rows.append({headers[i]: _cell_text(values_row[i] if i < len(values_row) else None) for i in range(len(headers))})
        finally:
            workbook.close()
    else:
        raise ValueError("Upload a CSV or XLSX file.")
    if not rows:
        raise ValueError("The file has a header row but no lead rows.")
    if len(rows) > 100000:
        raise ValueError("A single job can contain up to 100,000 rows. Split larger files into batches.")
    email_field = detect_email_field(headers)
    if not email_field:
        raise ValueError("No email column found. Name a column Email, Email Address, Work Email, or Business Email.")
    return headers, rows, email_field
